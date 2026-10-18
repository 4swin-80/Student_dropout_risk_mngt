from django.contrib.auth.models import User
from django.contrib.auth.hashers import check_password
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from tempfile import TemporaryDirectory


from .models import (
    Attendance,
    Department,
    Mentor,
    Student,
    StudentSupport,
    Teacher,
    UserProfile,
)

import io
import json
from datetime import date, timedelta
from decimal import Decimal
import openpyxl
from django.utils import timezone

from .models import (
    Attendance,
    AcademicPerformance,
    Assignment,
    ExamResult,
    RiskPrediction,
    Intervention,
    StudentSupportFollowUp,
    StudentNotification,
)
from .reports_utils import (
    build_full_report_data,
    sanitize_csv_cell,
)


class StudentNotificationTests(TestCase):
    def setUp(self):
        self.student_user = User.objects.create_user(
            username="notification-student", password="test-pass"
        )
        UserProfile.objects.create(user=self.student_user, role="STUDENT")
        self.student = Student.objects.create(
            student_id="NOTICE-1", user=self.student_user,
            first_name="Test", last_name="Student",
        )
        self.other_user = User.objects.create_user(
            username="notification-other", password="test-pass"
        )
        UserProfile.objects.create(user=self.other_user, role="STUDENT")
        self.other_student = Student.objects.create(
            student_id="NOTICE-2", user=self.other_user,
            first_name="Other", last_name="Student",
        )
        self.notification = StudentNotification.objects.create(
            student=self.student, title="Follow-up", message="Please meet your mentor.",
            category="Meeting",
        )
        self.private_notification = StudentNotification.objects.create(
            student=self.other_student, title="Private message", message="Secret content",
        )
        self.url = reverse("student_notifications")

    def test_student_sees_only_own_notifications_and_correct_counts(self):
        self.client.force_login(self.student_user)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Please meet your mentor.")
        self.assertNotContains(response, "Secret content")
        self.assertEqual(response.context["total_count"], 1)
        self.assertEqual(response.context["unread_count"], 1)
        self.assertEqual(response.context["read_count"], 0)

    def test_read_action_is_post_only_scoped_idempotent_and_updates_timestamp(self):
        self.client.force_login(self.student_user)
        action_url = reverse(
            "student_notification_mark_read", args=[self.notification.pk]
        )
        self.assertEqual(self.client.get(action_url).status_code, 405)
        self.notification.refresh_from_db()
        self.assertFalse(self.notification.is_read)
        self.assertEqual(self.client.post(action_url).status_code, 302)
        self.notification.refresh_from_db()
        self.assertTrue(self.notification.is_read)
        self.assertIsNotNone(self.notification.read_at)
        self.assertEqual(self.client.post(action_url).status_code, 302)
        foreign_url = reverse(
            "student_notification_mark_read", args=[self.private_notification.pk]
        )
        self.assertEqual(self.client.post(foreign_url).status_code, 404)
        self.private_notification.refresh_from_db()
        self.assertFalse(self.private_notification.is_read)

    def test_access_control_and_unlinked_student(self):
        self.assertEqual(self.client.get(self.url).status_code, 302)
        teacher = User.objects.create_user(username="notice-teacher", password="x")
        UserProfile.objects.create(user=teacher, role="TEACHER")
        self.client.force_login(teacher)
        self.assertEqual(self.client.get(self.url).status_code, 302)
        unlinked = User.objects.create_user(username="notice-unlinked", password="x")
        UserProfile.objects.create(user=unlinked, role="STUDENT")
        self.client.force_login(unlinked)
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_empty_and_all_read_states_and_safe_escaping(self):
        StudentNotification.objects.filter(student=self.student).delete()
        self.client.force_login(self.student_user)
        response = self.client.get(self.url)
        self.assertContains(response, "No Support Notifications Yet")
        self.assertEqual(response.context["total_count"], 0)
        read_notice = StudentNotification.objects.create(
            student=self.student, title="<b>Safe</b>", message="<script>unsafe()</script>",
            is_read=True,
        )
        response = self.client.get(self.url)
        self.assertEqual(response.context["total_count"], 1)
        self.assertEqual(response.context["unread_count"], 0)
        self.assertEqual(response.context["read_count"], 1)
        self.assertContains(response, "&lt;b&gt;Safe&lt;/b&gt;")
        self.assertNotContains(response, "<script>unsafe()</script>")

    def test_admin_can_create_explicit_student_safe_notification(self):
        admin = User.objects.create_user(username="notice-admin", password="x")
        UserProfile.objects.create(user=admin, role="ADMIN")
        self.client.force_login(admin)
        response = self.client.post(
            reverse("admin_student_notification_create"),
            {"student": self.student.pk, "title": "Meeting", "message": "Please meet us.", "category": "Meeting"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(StudentNotification.objects.filter(
            student=self.student, title="Meeting", created_by=admin
        ).exists())


class StudentAccountProvisioningTests(TestCase):
    def setUp(self):
        self.department = Department.objects.create(name="Computing")
        self.admin = User.objects.create_user(username="admin", password="admin-pass")
        UserProfile.objects.create(user=self.admin, role="ADMIN")
        mentor_user = User.objects.create_user(username="mentor", password="mentor-pass")
        UserProfile.objects.create(user=mentor_user, role="MENTOR")
        self.mentor = Mentor.objects.create(
            user=mentor_user, employee_id="M-TEST", department=self.department
        )

    def payload(self, student_id="STU-TEST", password=""):
        return {
            "student_id": student_id, "first_name": "Test", "last_name": "Student",
            "date_of_birth": "2005-01-01", "gender": "Other",
            "email": "test.student@example.com", "phone": "1234567890", "address": "",
            "department": str(self.department.pk), "course": "Computing", "semester": "1",
            "mentor": str(self.mentor.pk), "admission_date": "2024-01-01",
            "student_password": password,
        }

    def login_admin(self):
        self.client.force_login(self.admin)

    def test_password_field_is_rendered_for_add_and_unlinked_edit(self):
        self.login_admin()
        self.assertContains(self.client.get(reverse("student_create")), 'name="student_password"')
        student = Student.objects.create(student_id="STU-EDIT", first_name="Test")
        self.assertContains(
            self.client.get(reverse("student_update", args=[student.pk])),
            'name="student_password"',
        )

    def test_create_provisions_hashed_student_account(self):
        self.login_admin()
        response = self.client.post(reverse("student_create"), self.payload(password="ValidStrongPass!234"))
        self.assertEqual(response.status_code, 302)
        student = Student.objects.get(student_id="STU-TEST")
        self.assertEqual(student.user.username, student.student_id)
        self.assertEqual(student.user.profile.role, "STUDENT")
        self.assertTrue(student.user.check_password("ValidStrongPass!234"))

    def test_unlinked_blank_password_does_not_provision(self):
        self.login_admin()
        student = Student.objects.create(student_id="STU-EDIT", first_name="Old")
        response = self.client.post(reverse("student_update", args=[student.pk]), self.payload("STU-EDIT"))
        self.assertEqual(response.status_code, 302)
        student.refresh_from_db()
        self.assertIsNone(student.user_id)

    def test_safe_matching_account_is_reused_without_password_reset(self):
        self.login_admin()
        user = User.objects.create_user(
            username="STU-EDIT", password="OriginalPassword!234", first_name="Test",
            last_name="Student", email="test.student@example.com",
        )
        UserProfile.objects.create(user=user, role="STUDENT")
        student = Student.objects.create(student_id="STU-EDIT", first_name="Old")
        response = self.client.post(reverse("student_update", args=[student.pk]), self.payload("STU-EDIT"))
        self.assertEqual(response.status_code, 302)
        student.refresh_from_db()
        self.assertEqual(student.user_id, user.pk)
        self.assertTrue(user.check_password("OriginalPassword!234"))
        self.assertEqual(User.objects.filter(username="STU-EDIT").count(), 1)

    def test_incompatible_role_is_rejected_without_conversion(self):
        self.login_admin()
        user = User.objects.create_user(
            username="STU-EDIT", password="OriginalPassword!234", first_name="Test",
            last_name="Student", email="test.student@example.com",
        )
        UserProfile.objects.create(user=user, role="ADMIN")
        student = Student.objects.create(student_id="STU-EDIT", first_name="Old")
        response = self.client.post(reverse("student_update", args=[student.pk]), self.payload("STU-EDIT"))
        self.assertEqual(response.status_code, 200)
        student.refresh_from_db()
        user.profile.refresh_from_db()
        self.assertIsNone(student.user_id)
        self.assertEqual(user.profile.role, "ADMIN")


class TeacherStudentSupportReferralTests(TestCase):
    def setUp(self):
        self.department = Department.objects.create(
            name="Computer Science",
        )

        self.teacher_user = User.objects.create_user(
            username="teacher",
            password="password",
        )
        UserProfile.objects.create(
            user=self.teacher_user,
            role="TEACHER",
        )
        self.teacher = Teacher.objects.create(
            user=self.teacher_user,
            employee_id="T001",
            department=self.department,
        )

        self.mentor_user = User.objects.create_user(
            username="mentor",
            password="password",
        )
        UserProfile.objects.create(
            user=self.mentor_user,
            role="MENTOR",
        )
        self.mentor = Mentor.objects.create(
            user=self.mentor_user,
            employee_id="M001",
            department=self.department,
        )

        self.student = Student.objects.create(
            student_id="S001",
            first_name="Asha",
            last_name="Kumar",
            department=self.department,
        )

        StudentSupport.objects.create(
            student=self.student,
            teacher=self.teacher,
            mentor=self.mentor,
            support_type="Remedial teaching",
            reason="Existing active referral",
            status="Pending",
        )

    def test_duplicate_referral_warning_stays_on_student_report(self):
        self.client.login(
            username="teacher",
            password="password",
        )

        response = self.client.get(
            reverse(
                "teacher_refer_student_to_mentor",
                kwargs={"student_id": self.student.pk},
            ),
            follow=True,
        )

        self.assertRedirects(
            response,
            reverse(
                "teacher_student_support_report",
                kwargs={"student_id": self.student.pk},
            ),
        )
        self.assertTemplateUsed(
            response,
            "teacher/student_support_report.html",
        )
        self.assertContains(
            response,
            "This student already has an active support referral.",
        )
        self.assertEqual(
            StudentSupport.objects.filter(student=self.student).count(),
            1,
        )


# ============================================================
# ADMINISTRATOR - REPORTS & ANALYTICS TEST SUITE
# ============================================================


class AdminReportsAccessControlTests(TestCase):
    """
    Verify application-role authorization for all Reports & Analytics endpoints:
    - Page view: /admin/reports/
    - JSON API: /admin/reports/data/
    - CSV Export: /admin/reports/export/csv/
    - Excel Export: /admin/reports/export/excel/
    """

    def setUp(self):
        self.department = Department.objects.create(name="Computer Science")

        # Admin user (Application role: ADMIN)
        self.admin_user = User.objects.create_user(username="admin_user", password="password")
        UserProfile.objects.create(user=self.admin_user, role="ADMIN")

        # Teacher user (Application role: TEACHER)
        self.teacher_user = User.objects.create_user(username="teacher_user", password="password")
        UserProfile.objects.create(user=self.teacher_user, role="TEACHER")

        # Mentor user (Application role: MENTOR)
        self.mentor_user = User.objects.create_user(username="mentor_user", password="password")
        UserProfile.objects.create(user=self.mentor_user, role="MENTOR")

        # Superuser account (Explicitly restricted from application dashboards by policy)
        self.superuser = User.objects.create_superuser(username="super_user", password="password")
        UserProfile.objects.create(user=self.superuser, role="NONE")

        self.endpoints = [
            reverse("admin_reports"),
            reverse("admin_reports_data"),
            reverse("admin_reports_export_csv"),
            reverse("admin_reports_export_excel"),
        ]

    def test_unauthenticated_requests_redirect_to_login(self):
        for endpoint in self.endpoints:
            response = self.client.get(endpoint)
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response.url.startswith(reverse("login")))

    def test_teacher_denied_access_to_all_endpoints(self):
        self.client.login(username="teacher_user", password="password")
        for endpoint in self.endpoints:
            response = self.client.get(endpoint)
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response.url.startswith(reverse("login")))

    def test_mentor_denied_access_to_all_endpoints(self):
        self.client.login(username="mentor_user", password="password")
        for endpoint in self.endpoints:
            response = self.client.get(endpoint)
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response.url.startswith(reverse("login")))

    def test_superuser_denied_access_by_application_dashboard_policy(self):
        self.client.login(username="super_user", password="password")
        for endpoint in self.endpoints:
            response = self.client.get(endpoint)
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response.url.startswith(reverse("login")))

    def test_admin_granted_access_to_all_endpoints(self):
        self.client.login(username="admin_user", password="password")
        for endpoint in self.endpoints:
            response = self.client.get(endpoint)
            self.assertEqual(response.status_code, 200)


class AdminReportsFilterAndMetricsTests(TestCase):
    """
    Verify filtering logic and metric calculations:
    - Latest prediction per student without double-counting
    - Accurate unpredicted student tracking
    - Non-zero class attendance aggregation
    - 20-mark scale assignments with null marks handling
    - Overdue follow-up scheduling logic
    - Zero-denominator safety
    """

    def setUp(self):
        self.dep_cs = Department.objects.create(name="Computer Science")
        self.dep_ee = Department.objects.create(name="Electrical Engineering")

        self.teacher_user = User.objects.create_user(username="t_user", password="password")
        UserProfile.objects.create(user=self.teacher_user, role="TEACHER")
        self.teacher = Teacher.objects.create(user=self.teacher_user, employee_id="T10", department=self.dep_cs)

        self.mentor_user = User.objects.create_user(username="m_user", password="password")
        UserProfile.objects.create(user=self.mentor_user, role="MENTOR")
        self.mentor = Mentor.objects.create(user=self.mentor_user, employee_id="M10", department=self.dep_cs)

        # Student 1: CS, BCA, Sem 1
        self.s1 = Student.objects.create(
            student_id="CS01",
            first_name="Alice",
            last_name="Smith",
            department=self.dep_cs,
            course="BCA",
            semester=1,
            mentor=self.mentor,
        )

        # Student 2: CS, BCA, Sem 2
        self.s2 = Student.objects.create(
            student_id="CS02",
            first_name="Bob",
            last_name="Jones",
            department=self.dep_cs,
            course="BCA",
            semester=2,
            mentor=self.mentor,
        )

        # Student 3: EE, B.Tech, Sem 1 (Unpredicted)
        self.s3 = Student.objects.create(
            student_id="EE01",
            first_name="Charlie",
            last_name="Brown",
            department=self.dep_ee,
            course="B.Tech",
            semester=1,
            mentor=self.mentor,
        )

        # Risk Predictions:
        # S1 has an older Low prediction and a newer High prediction
        RiskPrediction.objects.create(
            student=self.s1,
            risk_score=Decimal("35.00"),
            risk_level="Low",
        )
        RiskPrediction.objects.create(
            student=self.s1,
            risk_score=Decimal("75.00"),
            risk_level="High",
        )
        # S2 has one Very High prediction
        RiskPrediction.objects.create(
            student=self.s2,
            risk_score=Decimal("85.00"),
            risk_level="Very High",
        )
        # S3 has NO prediction

        # Attendance:
        # S1: 10 attended / 10 conducted = 100%
        Attendance.objects.create(
            student=self.s1,
            date=date.today(),
            classes_conducted=10,
            classes_attended=10,
        )
        # S2: 6 attended / 10 conducted = 60%
        Attendance.objects.create(
            student=self.s2,
            date=date.today(),
            classes_conducted=10,
            classes_attended=6,
        )

        # Academic:
        AcademicPerformance.objects.create(
            student=self.s1,
            subject="Python",
            internal_marks=Decimal("80.00"),
            total_internal_marks=Decimal("100.00"),
        )
        ExamResult.objects.create(
            student=self.s1,
            exam_name="Midterm",
            subject="Python",
            marks_obtained=Decimal("70.00"),
            total_marks=Decimal("100.00"),
        )

        # Assignments:
        # One graded out of 20, one pending with marks=None
        Assignment.objects.create(
            student=self.s1,
            title="Lab 1",
            subject="Python",
            status="Submitted",
            marks=Decimal("18.00"),
        )
        Assignment.objects.create(
            student=self.s2,
            title="Lab 2",
            subject="Python",
            status="Pending",
            marks=None,
        )

        # Interventions:
        Intervention.objects.create(
            student=self.s1,
            mentor=self.mentor,
            intervention_type="Counseling",
            status="In Progress",
            intervention_date=date.today(),
            notes="Regular counseling",
        )

        # Referrals & Follow-ups:
        self.referral = StudentSupport.objects.create(
            student=self.s1,
            teacher=self.teacher,
            mentor=self.mentor,
            support_type="Remedial teaching",
            reason="Low quiz score",
            status="In Progress",
        )
        # Overdue follow-up (scheduled date in past, status=Pending)
        StudentSupportFollowUp.objects.create(
            referral=self.referral,
            scheduled_date=date.today() - timedelta(days=3),
            status="Pending",
        )
        # Completed follow-up
        StudentSupportFollowUp.objects.create(
            referral=self.referral,
            scheduled_date=date.today() - timedelta(days=1),
            status="Completed",
        )

    def test_current_risk_deduplication_and_unpredicted(self):
        data = build_full_report_data()
        summary = data["summary"]
        risk_dist = data["risk_distribution"]

        # Total distinct students
        self.assertEqual(summary["total_students"], 3)
        # S1 must only be counted once (latest: High); S2 is Very High
        self.assertEqual(summary["high_risk_count"], 1)
        self.assertEqual(summary["very_high_risk_count"], 1)
        self.assertEqual(summary["high_or_very_high_count"], 2)

        # S3 has no prediction, so unpredicted count is 1, not Very Low
        self.assertEqual(risk_dist["No Prediction"], 1)
        self.assertEqual(risk_dist["High"], 1)
        self.assertEqual(risk_dist["Very High"], 1)
        self.assertEqual(risk_dist["Low"], 0)  # S1's older Low prediction is superseded

        # Average risk score of predicted students: (75 + 85) / 2 = 80.0
        self.assertEqual(summary["average_risk_score"], 80.0)

    def test_filter_by_department(self):
        data = build_full_report_data(department_id=self.dep_cs.id)
        self.assertEqual(data["summary"]["total_students"], 2)
        self.assertEqual(data["filters"]["active"]["department_id"], self.dep_cs.id)

    def test_filter_by_course_and_semester(self):
        data = build_full_report_data(
            department_id=self.dep_cs.id,
            course="BCA",
            semester=1,
        )
        self.assertEqual(data["summary"]["total_students"], 1)
        self.assertEqual(data["summary"]["high_risk_count"], 1)

    def test_attendance_aggregation_formula(self):
        data = build_full_report_data(department_id=self.dep_cs.id)
        # Total conducted: 10 + 10 = 20, attended: 10 + 6 = 16 => 80.0%
        att = data["attendance"]
        self.assertEqual(att["total_conducted"], 20)
        self.assertEqual(att["total_attended"], 16)
        self.assertEqual(att["overall_percentage"], 80.0)

    def test_assignment_statistics_respect_20_mark_scale(self):
        data = build_full_report_data()
        acad = data["academics"]
        self.assertEqual(acad["assignment_total_count"], 2)
        self.assertEqual(acad["assignment_graded_count"], 1)
        self.assertEqual(acad["assignment_average_marks_20"], 18.0)
        self.assertEqual(acad["assignment_status"]["Submitted"], 1)
        self.assertEqual(acad["assignment_status"]["Pending"], 1)

    def test_overdue_follow_up_detection(self):
        data = build_full_report_data()
        supp = data["support_and_follow_ups"]
        self.assertEqual(supp["total_follow_ups"], 2)
        self.assertEqual(supp["completed_follow_ups"], 1)
        self.assertEqual(supp["pending_follow_ups"], 1)
        self.assertEqual(supp["overdue_follow_ups"], 1)

    def test_zero_denominator_safe_handling(self):
        dep_empty = Department.objects.create(name="Empty Dept")
        s_new = Student.objects.create(
            student_id="NEW01",
            first_name="Eve",
            department=dep_empty,
        )
        Attendance.objects.create(
            student=s_new,
            date=date.today(),
            classes_conducted=0,
            classes_attended=0,
        )
        data = build_full_report_data(department_id=dep_empty.id)
        # Should cleanly return None, not crash or 0
        self.assertIsNone(data["summary"]["average_attendance_percentage"])
        self.assertIsNone(data["attendance"]["overall_percentage"])


class AdminReportsExportTests(TestCase):
    """
    Verify CSV and Excel export formats, privacy restrictions, and formula sanitization.
    """

    def setUp(self):
        self.admin_user = User.objects.create_user(username="admin_exporter", password="password")
        UserProfile.objects.create(user=self.admin_user, role="ADMIN")
        self.client.login(username="admin_exporter", password="password")

        self.dep = Department.objects.create(name="Data Science")
        self.student = Student.objects.create(
            student_id="SECURE01",
            first_name="ConfidentialFirst",
            last_name="ConfidentialLast",
            email="confidential@example.com",
            phone="9998887777",
            department=self.dep,
        )
        RiskPrediction.objects.create(
            student=self.student,
            risk_score=Decimal("65.00"),
            risk_level="High",
        )

    def test_csv_export_structure_and_privacy(self):
        response = self.client.get(reverse("admin_reports_export_csv"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue("text/csv" in response["Content-Type"])
        self.assertTrue("attachment" in response["Content-Disposition"])

        content = response.content.decode("utf-8")
        # Check headings
        self.assertIn("EXECUTIVE SUMMARY", content)
        self.assertIn("CURRENT RISK DISTRIBUTION", content)
        self.assertIn("DEPARTMENT COMPARISONS", content)

        # Privacy verification: No individual student identifiers must leak in export
        self.assertNotIn("SECURE01", content)
        self.assertNotIn("ConfidentialFirst", content)
        self.assertNotIn("ConfidentialLast", content)
        self.assertNotIn("confidential@example.com", content)
        self.assertNotIn("9998887777", content)

    def test_csv_formula_injection_sanitization(self):
        self.assertEqual(sanitize_csv_cell("=SUM(A1:A10)"), "'=SUM(A1:A10)")
        self.assertEqual(sanitize_csv_cell("+12345"), "'+12345")
        self.assertEqual(sanitize_csv_cell("-CmdPrompt"), "'-CmdPrompt")
        self.assertEqual(sanitize_csv_cell("@ExternalLink"), "'@ExternalLink")
        self.assertEqual(sanitize_csv_cell("Normal Text"), "Normal Text")

    def test_excel_export_structure_and_sheets(self):
        response = self.client.get(reverse("admin_reports_export_excel"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue("openxmlformats" in response["Content-Type"])

        # Load workbook from response bytes
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        expected_sheets = [
            "Overview",
            "Risk Analytics",
            "Attendance & Academics",
            "Assignments",
            "Interventions & Follow-ups",
            "Department Comparisons",
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Verify no individual student identifiers in any sheet
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            for row in ws.iter_rows(values_only=True):
                for cell_val in row:
                    val_str = str(cell_val or "")
                    self.assertNotIn("SECURE01", val_str)
                    self.assertNotIn("ConfidentialFirst", val_str)
                    self.assertNotIn("ConfidentialLast", val_str)



# ============================================================
# STUDENT DASHBOARD - ACCESS CONTROL AND DATA ISOLATION
# ============================================================


class StudentDashboardAccessTests(TestCase):

    def setUp(self):
        self.department = Department.objects.create(
            name="Dashboard Test Department"
        )

        # Create the first student and login account.
        self.student_user = User.objects.create_user(
            username="DASH-STU-001",
            password="StudentPass!123",
        )

        UserProfile.objects.create(
            user=self.student_user,
            role="STUDENT",
        )

        self.student = Student.objects.create(
            student_id="DASH-STU-001",
            user=self.student_user,
            first_name="Anu",
            last_name="Student",
            department=self.department,
        )

        # Create another student to verify data isolation.
        self.other_student_user = User.objects.create_user(
            username="DASH-STU-002",
            password="StudentPass!456",
        )

        UserProfile.objects.create(
            user=self.other_student_user,
            role="STUDENT",
        )

        self.other_student = Student.objects.create(
            student_id="DASH-STU-002",
            user=self.other_student_user,
            first_name="Other",
            last_name="Student",
            department=self.department,
        )

        # Create users belonging to other roles.
        self.teacher_user = User.objects.create_user(
            username="DASH-TEACHER",
            password="TeacherPass!123",
        )

        UserProfile.objects.create(
            user=self.teacher_user,
            role="TEACHER",
        )

        self.mentor_user = User.objects.create_user(
            username="DASH-MENTOR",
            password="MentorPass!123",
        )

        UserProfile.objects.create(
            user=self.mentor_user,
            role="MENTOR",
        )

        # Create a superuser to verify that superuser status
        # alone does not grant Student Dashboard access.
        self.superuser = User.objects.create_superuser(
            username="DASH-SUPERUSER",
            email="dashboard-admin@example.com",
            password="SuperPass!123",
        )

    def test_unauthenticated_user_is_redirected_to_login(self):
        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertEqual(response.status_code, 302)
        self.assertIn(
            reverse("login"),
            response.url,
        )

    def test_linked_student_can_access_dashboard(self):
        self.client.force_login(self.student_user)

        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.context["student"],
            self.student,
        )

    def test_linked_student_can_access_profile_page(self):
        self.client.force_login(self.student_user)

        response = self.client.get(reverse("student_profile"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["student"], self.student)
        self.assertContains(response, self.student.student_id)
        self.assertContains(response, self.student.first_name)
        self.assertContains(response, "My Profile")
        self.assertEqual(reverse("student_profile"), "/student-profile/")

    def test_profile_query_parameter_cannot_switch_student(self):
        self.client.force_login(self.student_user)

        response = self.client.get(
            reverse("student_profile"),
            {"student_id": self.other_student.student_id},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["student"], self.student)
        self.assertNotContains(response, self.other_student.student_id)

    def test_unauthenticated_user_is_redirected_from_profile(self):
        response = self.client.get(reverse("student_profile"))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

    def test_incompatible_role_cannot_access_profile(self):
        self.client.force_login(self.teacher_user)

        response = self.client.get(reverse("student_profile"))

        self.assertIn(response.status_code, [302, 403])

    def test_unlinked_student_profile_is_forbidden_gracefully(self):
        unlinked_user = User.objects.create_user(
            username="PROFILE-UNLINKED",
            password="UnlinkedPass!123",
        )
        UserProfile.objects.create(user=unlinked_user, role="STUDENT")
        self.client.force_login(unlinked_user)

        response = self.client.get(reverse("student_profile"))

        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "not linked to this account", status_code=403)

    def test_student_can_update_personal_profile_fields(self):
        self.client.force_login(self.student_user)

        response = self.client.post(
            reverse("student_profile"),
            {
                "first_name": "Anushka",
                "last_name": "Student",
                "email": "anushka@example.com",
                "phone": "9876543210",
                "date_of_birth": "2004-05-06",
                "gender": "Female",
                "address": "12 College Road",
            },
        )

        self.assertRedirects(response, reverse("student_profile"))
        self.student.refresh_from_db()
        self.assertEqual(self.student.first_name, "Anushka")
        self.assertEqual(self.student.last_name, "Student")
        self.assertEqual(self.student.email, "anushka@example.com")
        self.assertEqual(self.student.phone, "9876543210")
        self.assertEqual(self.student.date_of_birth, date(2004, 5, 6))
        self.assertEqual(self.student.gender, "Female")
        self.assertEqual(self.student.address, "12 College Road")

    def test_student_can_upload_and_retrieve_own_profile_photo(self):
        # A valid 1x1 PNG image.
        import base64

        png_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/"
            "pS8AAAAASUVORK5CYII="
        )
        self.client.force_login(self.student_user)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.client.post(
                reverse("student_profile"),
                {
                    "first_name": self.student.first_name,
                    "last_name": self.student.last_name,
                    "email": self.student.email,
                    "phone": self.student.phone,
                    "date_of_birth": "",
                    "gender": "",
                    "address": "",
                    "profile_photo": SimpleUploadedFile(
                        "profile.png", png_bytes, content_type="image/png"
                    ),
                },
            )
            self.assertRedirects(response, reverse("student_profile"))
            self.student.refresh_from_db()

            photo_response = self.client.get(reverse("student_profile_photo"))
            self.assertEqual(photo_response.status_code, 200)
            self.assertEqual(photo_response["Content-Type"], "image/png")
            photo_response.close()

    def test_invalid_profile_photo_is_rejected(self):
        self.client.force_login(self.student_user)

        response = self.client.post(
            reverse("student_profile"),
            {
                "first_name": self.student.first_name,
                "last_name": self.student.last_name,
                "email": self.student.email,
                "phone": self.student.phone,
                "date_of_birth": "",
                "gender": "",
                "address": "",
                "profile_photo": SimpleUploadedFile(
                    "not-an-image.png", b"not an image", content_type="image/png"
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Upload a valid JPEG, PNG, GIF, or WebP image.")
        self.student.refresh_from_db()
        self.assertFalse(self.student.profile_photo)

    def test_non_student_cannot_retrieve_student_profile_photo(self):
        self.client.force_login(self.teacher_user)

        response = self.client.get(reverse("student_profile_photo"))

        self.assertIn(response.status_code, [302, 403])

    def test_student_cannot_view_another_students_profile(self):
        self.client.force_login(self.student_user)

        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.context["student"],
            self.student,
        )
        self.assertNotEqual(
            response.context["student"],
            self.other_student,
        )

    def test_query_parameter_cannot_switch_student(self):
        self.client.force_login(self.student_user)

        response = self.client.get(
            reverse("student_dashboard"),
            {"student_id": self.other_student.student_id},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.context["student"],
            self.student,
        )

    def test_student_without_linked_record_is_forbidden(self):
        unlinked_user = User.objects.create_user(
            username="DASH-UNLINKED",
            password="UnlinkedPass!123",
        )

        UserProfile.objects.create(
            user=unlinked_user,
            role="STUDENT",
        )

        self.client.force_login(unlinked_user)

        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertEqual(response.status_code, 403)

    def test_teacher_cannot_access_student_dashboard(self):
        self.client.force_login(self.teacher_user)

        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertIn(response.status_code, [302, 403])

    def test_mentor_cannot_access_student_dashboard(self):
        self.client.force_login(self.mentor_user)

        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertIn(response.status_code, [302, 403])

    def test_superuser_cannot_access_student_dashboard(self):
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse("student_dashboard")
        )

        self.assertIn(response.status_code, [302, 403])


class StudentAssignmentsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="assign-student", password="pass")
        UserProfile.objects.create(user=self.user, role="STUDENT")
        self.student = Student.objects.create(
            student_id="ASSIGN-001", user=self.user, first_name="Asha"
        )
        self.other_user = User.objects.create_user(username="assign-other", password="pass")
        UserProfile.objects.create(user=self.other_user, role="STUDENT")
        self.other_student = Student.objects.create(
            student_id="ASSIGN-002", user=self.other_user, first_name="Other"
        )
        self.client.force_login(self.user)

    def test_student_assignments_url_and_authenticated_page(self):
        self.assertEqual(reverse("student_assignments"), "/student-assignments/")
        response = self.client.get(reverse("student_assignments"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "My Assignments")
        self.assertContains(response, f'href="{reverse("student_dashboard")}"')

    def test_assignment_data_and_summary_are_isolated(self):
        Assignment.objects.create(
            student=self.student,
            title="My Project", subject="Science",
            submission_date=date(2026, 5, 2), status="Submitted",
            marks=Decimal("16.00"),
        )
        Assignment.objects.create(
            student=self.student, title="Draft", subject="Math",
            status="Pending", marks=None,
        )
        Assignment.objects.create(
            student=self.student, title="Old Work", subject="History",
            status="Late", marks=Decimal("21.00"),
        )
        Assignment.objects.create(
            student=self.other_student, title="Private Assignment",
            subject="Secret", status="Late",
        )

        response = self.client.get(
            reverse("student_assignments"),
            {"student_id": self.other_student.student_id},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["assignment_count"], 3)
        self.assertEqual(response.context["pending_assignment_count"], 1)
        self.assertEqual(response.context["submitted_assignment_count"], 1)
        self.assertEqual(response.context["late_assignment_count"], 1)
        self.assertContains(response, "My Project")
        self.assertContains(response, "Science")
        self.assertContains(response, "02 May 2026")
        self.assertContains(response, "16.00 / 20")
        self.assertContains(response, "Not specified")
        self.assertContains(response, "Not graded")
        self.assertNotContains(response, "Private Assignment")
        self.assertNotContains(response, "Secret")

    def test_empty_state_and_zero_counts(self):
        response = self.client.get(reverse("student_assignments"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "You don't have any assignments recorded yet.")
        self.assertEqual(response.context["assignment_count"], 0)
        self.assertEqual(response.context["pending_assignment_count"], 0)
        self.assertEqual(response.context["submitted_assignment_count"], 0)
        self.assertEqual(response.context["late_assignment_count"], 0)

    def test_unauthenticated_and_incompatible_roles_are_rejected(self):
        self.client.logout()
        response = self.client.get(reverse("student_assignments"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

        for role in ("TEACHER", "MENTOR", "ADMIN"):
            user = User.objects.create_user(username=f"assign-{role.lower()}", password="pass")
            UserProfile.objects.create(user=user, role=role)
            self.client.force_login(user)
            denied = self.client.get(reverse("student_assignments"))
            self.assertIn(denied.status_code, [302, 403])

        superuser = User.objects.create_superuser(
            username="assign-super", email="assign-super@example.com", password="pass"
        )
        self.client.force_login(superuser)
        self.assertIn(self.client.get(reverse("student_assignments")).status_code, [302, 403])

    def test_student_without_linked_record_is_forbidden(self):
        unlinked = User.objects.create_user(username="assign-unlinked", password="pass")
        UserProfile.objects.create(user=unlinked, role="STUDENT")
        self.client.force_login(unlinked)
        response = self.client.get(reverse("student_assignments"))
        self.assertEqual(response.status_code, 403)


class StudentExaminationResultsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="exam-student", password="pass")
        UserProfile.objects.create(user=self.user, role="STUDENT")
        self.student = Student.objects.create(
            student_id="EXAM-001", user=self.user, first_name="Asha"
        )
        self.other_user = User.objects.create_user(username="exam-other", password="pass")
        UserProfile.objects.create(user=self.other_user, role="STUDENT")
        self.other_student = Student.objects.create(
            student_id="EXAM-002", user=self.other_user, first_name="Other"
        )
        self.client.force_login(self.user)

    def create_result(self, student, exam_name, subject, marks, total, **kwargs):
        return ExamResult.objects.create(
            student=student, exam_name=exam_name, subject=subject,
            marks_obtained=Decimal(marks), total_marks=Decimal(total), **kwargs,
        )

    def test_url_and_page_summary_are_student_scoped(self):
        self.assertEqual(
            reverse("student_examination_results"), "/student-examination-results/"
        )
        self.create_result(self.student, "Midterm", "Mathematics", "40", "50")
        self.create_result(self.student, "Final", "Science", "90", "100")
        self.create_result(self.other_student, "Private Exam", "Secret", "100", "100")

        response = self.client.get(
            reverse("student_examination_results"),
            {"student_id": self.other_student.student_id},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["exam_record_count"], 2)
        self.assertEqual(response.context["valid_result_count"], 2)
        self.assertEqual(response.context["awaiting_marks_count"], 0)
        self.assertEqual(response.context["exam_average"], Decimal("85.00"))
        self.assertContains(response, "Mathematics")
        self.assertContains(response, "40.00")
        self.assertContains(response, "80.00%")
        self.assertNotContains(response, "Private Exam")
        self.assertNotContains(response, "Secret")

    def test_invalid_results_are_excluded_and_never_show_percentages(self):
        self.create_result(self.student, "Good", "History", "75", "100")
        self.create_result(self.student, "Zero total", "Physics", "0", "0")
        self.create_result(self.student, "Too high", "Chemistry", "101", "100")
        self.create_result(self.student, "Negative", "Biology", "-1", "100")

        response = self.client.get(reverse("student_examination_results"))
        self.assertEqual(response.context["exam_record_count"], 4)
        self.assertEqual(response.context["valid_result_count"], 1)
        self.assertEqual(response.context["exam_average"], Decimal("75.00"))
        self.assertContains(response, "Zero total")
        self.assertContains(response, "Unavailable")
        self.assertNotContains(response, "101.00%")
        self.assertNotContains(response, "-1.00%")

    def test_empty_state_and_dashboard_link_and_average(self):
        response = self.client.get(reverse("student_examination_results"))
        self.assertContains(response, "No examination results have been recorded yet.")
        self.assertEqual(response.context["exam_record_count"], 0)
        self.assertEqual(response.context["valid_result_count"], 0)
        self.assertEqual(response.context["awaiting_marks_count"], 0)
        self.assertIsNone(response.context["exam_average"])

        dashboard = self.client.get(reverse("student_dashboard"))
        self.assertContains(dashboard, "View Examination Results")
        self.create_result(self.student, "Dashboard check", "Math", "-1", "100")
        self.create_result(self.student, "Dashboard valid", "Math", "50", "100")
        dashboard = self.client.get(reverse("student_dashboard"))
        self.assertEqual(dashboard.context["exam_count"], 2)
        self.assertEqual(dashboard.context["exam_average"], Decimal("50.00"))

    def test_missing_student_record_and_role_access(self):
        self.client.logout()
        response = self.client.get(reverse("student_examination_results"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

        for role in ("TEACHER", "MENTOR", "ADMIN"):
            user = User.objects.create_user(username=f"exam-{role.lower()}", password="pass")
            UserProfile.objects.create(user=user, role=role)
            self.client.force_login(user)
            self.assertIn(
                self.client.get(reverse("student_examination_results")).status_code,
                [302, 403],
            )

        superuser = User.objects.create_superuser(
            username="exam-super", email="exam-super@example.com", password="pass"
        )
        self.client.force_login(superuser)
        self.assertIn(
            self.client.get(reverse("student_examination_results")).status_code,
            [302, 403],
        )

        unlinked = User.objects.create_user(username="exam-unlinked", password="pass")
        UserProfile.objects.create(user=unlinked, role="STUDENT")
        self.client.force_login(unlinked)
        self.assertEqual(
            self.client.get(reverse("student_examination_results")).status_code, 403
        )


class StudentPerformanceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="performance-student", password="pass")
        UserProfile.objects.create(user=self.user, role="STUDENT")
        self.student = Student.objects.create(
            student_id="PERF-001", user=self.user, first_name="Asha"
        )
        other_user = User.objects.create_user(username="performance-other", password="pass")
        UserProfile.objects.create(user=other_user, role="STUDENT")
        self.other_student = Student.objects.create(
            student_id="PERF-002", user=other_user, first_name="Other"
        )
        self.client.force_login(self.user)

    def test_student_performance_page_has_private_chart_data_and_safe_serialization(self):
        self.assertEqual(reverse("student_performance"), "/student-performance/")
        AcademicPerformance.objects.create(
            student=self.student, subject="Math <script>", internal_marks=0,
            total_internal_marks=50,
        )
        AcademicPerformance.objects.create(
            student=self.student, subject="Math <script>", internal_marks=30,
            total_internal_marks=50,
        )
        AcademicPerformance.objects.create(
            student=self.student, subject="Invalid", internal_marks=80,
            total_internal_marks=0,
        )
        AcademicPerformance.objects.create(
            student=self.other_student, subject="Secret academic", internal_marks=100,
            total_internal_marks=100,
        )
        ExamResult.objects.create(
            student=self.student, exam_name="Midterm", subject="Science",
            marks_obtained=Decimal("40"), total_marks=Decimal("50"),
            exam_date=date(2025, 1, 10),
        )
        ExamResult.objects.create(
            student=self.student, exam_name="Midterm", subject="Science",
            marks_obtained=Decimal("45"), total_marks=Decimal("50"),
            exam_date=date(2025, 2, 10),
        )
        ExamResult.objects.create(
            student=self.other_student, exam_name="Secret Exam", subject="Secret",
            marks_obtained=Decimal("100"), total_marks=Decimal("100"),
        )

        response = self.client.get(reverse("student_performance"), {"student_id": self.other_student.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["academic_average"], Decimal("30.00"))
        self.assertEqual(response.context["exam_average"], Decimal("85.00"))
        self.assertEqual(response.context["valid_academic_record_count"], 2)
        self.assertEqual(response.context["valid_exam_record_count"], 2)
        self.assertEqual([p["percentage"] for p in response.context["academic_chart_data"]], [0.0, 60.0])
        self.assertEqual([p["date"] for p in response.context["performance_trend_data"]], ["2025-01-10", "2025-02-10"])
        self.assertNotContains(response, "Secret academic")
        self.assertNotContains(response, "Secret Exam")
        self.assertContains(response, "academic-chart-data")
        self.assertContains(response, "\\u003Cscript\\u003E")
        self.assertContains(response, f'href="{reverse("student_dashboard")}"')
        dashboard = self.client.get(reverse("student_dashboard"))
        self.assertContains(dashboard, f'href="{reverse("student_performance")}"')

    def test_empty_state_and_unavailable_trend(self):
        response = self.client.get(reverse("student_performance"))
        self.assertEqual(response.context["academic_chart_data"], [])
        self.assertEqual(response.context["exam_chart_data"], [])
        self.assertIsNone(response.context["academic_average"])
        self.assertIsNone(response.context["exam_average"])
        self.assertFalse(response.context["has_performance_trend"])
        self.assertContains(response, "No data available")

    def test_student_role_and_linked_student_required(self):
        unlinked = User.objects.create_user(username="performance-unlinked", password="pass")
        UserProfile.objects.create(user=unlinked, role="STUDENT")
        self.client.force_login(unlinked)
        self.assertEqual(self.client.get(reverse("student_performance")).status_code, 403)

        teacher = User.objects.create_user(username="performance-teacher", password="pass")
        UserProfile.objects.create(user=teacher, role="TEACHER")
        self.client.force_login(teacher)
        self.assertIn(self.client.get(reverse("student_performance")).status_code, [302, 403])

